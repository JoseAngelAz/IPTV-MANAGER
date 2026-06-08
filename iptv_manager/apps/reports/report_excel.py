from io import BytesIO
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, Reference


HEADER_FILL = PatternFill(start_color='1e293b', end_color='1e293b', fill_type='solid')
HEADER_FONT = Font(name='Calibri', bold=True, color='ffffff', size=11)
ALT_FILL = PatternFill(start_color='f0f4ff', end_color='f0f4ff', fill_type='solid')
BORDER = Border(
    left=Side(style='thin', color='cbd5e1'),
    right=Side(style='thin', color='cbd5e1'),
    top=Side(style='thin', color='cbd5e1'),
    bottom=Side(style='thin', color='cbd5e1'),
)


def _style_sheet(ws):
    ws.sheet_properties.tabColor = '2563eb'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def _write_header(ws, title, date_start, date_end):
    ws.merge_cells('A1:Z1')
    ws['A1'] = title
    ws['A1'].font = Font(name='Calibri', bold=True, size=16, color='2563eb')
    ws['A2'] = f'Generado: {datetime.now().strftime("%d/%m/%Y %H:%M")}'
    ws['A2'].font = Font(name='Calibri', size=9, color='64748b')
    if date_start and date_end:
        ws['B2'] = f'Período: {date_start} al {date_end}'
        ws['B2'].font = Font(name='Calibri', size=9, color='64748b')
    ws.row_dimensions[3].height = 8


def _write_table(ws, start_row, headers, rows):
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = BORDER

    for row_idx, row_data in enumerate(rows, start_row + 1):
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name='Calibri', size=10)
            cell.border = BORDER
            cell.alignment = Alignment(vertical='center')
            if (row_idx - start_row) % 2 == 0:
                cell.fill = ALT_FILL

    data_rows = len(rows)
    for col_idx in range(1, len(headers) + 1):
        max_len = len(str(headers[col_idx - 1]))
        for row_idx in range(start_row, start_row + data_rows + 1):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val:
                max_len = max(max_len, min(len(str(val)), 50))
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 3

    return start_row + data_rows + 1


def generate_report_excel(sections, data, date_start, date_end, user):
    wb = Workbook()
    first = True
    labels = {'clientes': 'Clientes', 'finanzas': 'Finanzas', 'notificaciones': 'Notificaciones', 'suscripciones': 'Suscripciones'}

    for key in sections:
        if first:
            ws = wb.active
            ws.title = labels.get(key, key)
            first = False
        else:
            ws = wb.create_sheet(title=labels.get(key, key))

        _style_sheet(ws)
        title = f'IPTV Manager — Reporte de {labels.get(key, key)}'
        _write_header(ws, title, date_start, date_end)

        if key == 'clientes':
            clients = data.get('clientes', [])
            headers = ['#', 'Nombre', 'Teléfono', 'Email', 'Dispositivo', 'Registro', 'Subs Activas', 'Total Gastado']
            rows = []
            for i, c in enumerate(clients, 1):
                disp = (c.get('dispositivo_id') or '').replace('-', '')
                fecha_raw = c.get('fecha_registro', '')
                if isinstance(fecha_raw, datetime):
                    fecha_str = fecha_raw.strftime('%Y-%m-%d')
                else:
                    fecha_str = str(fecha_raw)[:10] if fecha_raw else ''
                total = float(c.get('total_gastado') or 0)
                rows.append([i, c.get('nombre', ''), c.get('telefono', ''), c.get('email', ''),
                            disp, fecha_str,
                            c.get('subs_activas', 0) or 0,
                            f'${total:,.2f}'])
            _write_table(ws, 4, headers, rows)
            ws.cell(row=4 + len(rows) + 1, column=1,
                    value=f'Total clientes: {len(clients)}').font = Font(name='Calibri', italic=True, size=9, color='64748b')

        elif key == 'finanzas':
            resumen = data.get('resumen_finanzas', {})
            ws.cell(row=4, column=1, value='Resumen Financiero').font = Font(bold=True, size=11, color='2563eb')
            ws.cell(row=4, column=4, value='ChartData').font = Font(bold=True, size=9, color='94a3b8')
            ws.cell(row=5, column=4, value='Ingresos').font = Font(name='Calibri', size=10, bold=True)
            ws.cell(row=5, column=5, value=float(resumen.get('ingresos', 0))).font = Font(name='Calibri', size=10)
            ws.cell(row=6, column=4, value='Egresos').font = Font(name='Calibri', size=10, bold=True)
            ws.cell(row=6, column=5, value=float(resumen.get('egresos', 0))).font = Font(name='Calibri', size=10)
            r = 5
            for label, field in [('Ingresos', 'ingresos'), ('Egresos', 'egresos'), ('Neto', 'neto')]:
                val = float(resumen.get(field, 0))
                ws.cell(row=r, column=1, value=label).font = Font(name='Calibri', size=10, bold=True)
                cell = ws.cell(row=r, column=2, value=f'${val:,.2f}')
                cell.font = Font(name='Calibri', size=10,
                                 color='16a34a' if field != 'egresos' else 'dc2626')
                r += 1

            if float(resumen.get('ingresos', 0)) > 0 or float(resumen.get('egresos', 0)) > 0:
                chart = BarChart()
                chart.type = 'col'
                chart.title = 'Ingresos vs Egresos'
                chart.y_axis.title = 'Monto ($)'
                chart.style = 10
                chart.width = 14
                chart.height = 8
                data_ref = Reference(ws, min_col=5, min_row=5, max_row=6)
                cats_ref = Reference(ws, min_col=4, min_row=5, max_row=6)
                chart.add_data(data_ref, titles_from_data=False)
                chart.set_categories(cats_ref)
                ws.add_chart(chart, 'D8')

            finanzas = data.get('finanzas', [])
            headers = ['#', 'Tipo', 'Monto', 'Fecha', 'Descripción']
            rows = []
            for i, m in enumerate(finanzas, 1):
                fecha_raw = m.get('fecha', '')
                if isinstance(fecha_raw, datetime):
                    fecha_str = fecha_raw.strftime('%Y-%m-%d')
                else:
                    fecha_str = str(fecha_raw)[:10] if fecha_raw else ''
                rows.append([i, m.get('tipo', ''), f'${float(m.get("monto", 0)):,.2f}',
                            fecha_str, m.get('descripcion', '')[:50]])
            _write_table(ws, r + 1, headers, rows)

        elif key == 'notificaciones':
            notifs = data.get('notificaciones', [])
            headers = ['#', 'Canal', 'Destinatario', 'Cliente', 'Mensaje', 'Estado', 'Fecha']
            rows = []
            for i, n in enumerate(notifs, 1):
                fecha_raw = n.get('fecha_creacion', '')
                if isinstance(fecha_raw, datetime):
                    fecha_str = fecha_raw.strftime('%Y-%m-%d')
                else:
                    fecha_str = str(fecha_raw)[:10] if fecha_raw else ''
                rows.append([i, n.get('canal', ''), n.get('destinatario', ''), n.get('cliente_nombre', ''),
                            n.get('mensaje', '')[:50],
                            'Enviado' if n.get('enviado_ok') else 'Fallido',
                            fecha_str])
            _write_table(ws, 4, headers, rows)

        elif key == 'suscripciones':
            subs = data.get('suscripciones', [])
            if subs:
                estados_count = {}
                for s in subs:
                    est = s.get('estado', '')
                    estados_count[est] = estados_count.get(est, 0) + 1
                ws.cell(row=4, column=6, value='Estado').font = Font(bold=True, size=9, color='94a3b8')
                ws.cell(row=4, column=7, value='Cantidad').font = Font(bold=True, size=9, color='94a3b8')
                col = 6
                for i, (est, cnt) in enumerate(sorted(estados_count.items()), 5):
                    ws.cell(row=i, column=col, value=est).font = Font(name='Calibri', size=10, bold=True)
                    ws.cell(row=i, column=col + 1, value=cnt).font = Font(name='Calibri', size=10)
                if len(estados_count) > 0:
                    chart = BarChart()
                    chart.type = 'col'
                    chart.title = 'Suscripciones por Estado'
                    chart.y_axis.title = 'Cantidad'
                    chart.style = 10
                    chart.width = 14
                    chart.height = 8
                    data_ref = Reference(ws, min_col=7, min_row=5, max_row=4 + len(estados_count))
                    cats_ref = Reference(ws, min_col=6, min_row=5, max_row=4 + len(estados_count))
                    chart.add_data(data_ref, titles_from_data=False)
                    chart.set_categories(cats_ref)
                    ws.add_chart(chart, 'F8')

            headers = ['#', 'Cliente', 'Teléfono', 'Plan', 'Precio', 'Inicio', 'Vencimiento', 'Días Rest.', 'Estado', 'Método de pago']
            rows = []
            for i, s in enumerate(subs, 1):
                fecha_ini = s.get('fecha_inicio', '')
                fecha_ven = s.get('fecha_vencimiento', '')
                for fld in ('fecha_inicio', 'fecha_vencimiento'):
                    raw = s.get(fld, '')
                    if isinstance(raw, datetime):
                        s[fld] = raw.strftime('%Y-%m-%d')
                    else:
                        s[fld] = str(raw)[:10] if raw else ''
                rows.append([i, s.get('cliente_nombre', ''), s.get('cliente_telefono', ''),
                            s.get('plan_nombre', ''),
                            f'${float(s.get("plan_precio", 0)):,.2f}',
                            s['fecha_inicio'], s['fecha_vencimiento'],
                            s.get('dias_restantes', 0),
                            s.get('estado', ''),
                            s.get('metodo_pago', '')])
            _write_table(ws, 4, headers, rows)
            total_row = 4 + len(rows) + 1
            ws.cell(row=total_row, column=1,
                    value=f'Total suscripciones: {len(subs)}').font = Font(name='Calibri', italic=True, size=9, color='64748b')

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
